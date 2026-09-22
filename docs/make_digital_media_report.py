# -*- coding: utf-8 -*-
"""รายงานวิชาสื่อดิจิทัล — โครงหัวข้อที่แนะนำ"""
from __future__ import annotations

import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = os.path.dirname(__file__)
OUTS = [
    os.path.join(BASE, "รายงานสื่อดิจิทัล-โครงแนะนำ.docx"),
    os.path.join(os.path.expanduser("~"), "Downloads", "รายงานสื่อดิจิทัล-โครงแนะนำ.docx"),
]
SEAL = os.path.join(BASE, "assets", "ivec-seal-user.png")
SEAL_FB = os.path.join(BASE, "assets", "ivec3-seal.png")
FONT = "TH Sarabun New"
STUDENT = "นายพลพล กำมะถัน"
SID = "69419100009"
TEACHER = "นางชญานิตย์ ภาสว่าง"


def font(run, size=16, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), FONT)


def P(doc, text="", *, size=16, bold=False, align="left", before=0, after=6, first=0, indent=0, spacing=1.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = spacing
    if first:
        pf.first_line_indent = Cm(first)
    if indent:
        pf.left_indent = Cm(indent)
    p.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    if text:
        font(p.add_run(text), size=size, bold=bold)
    return p


def H1(doc, t):
    P(doc, t, size=20, bold=True, align="center", before=0, after=8)


def H2(doc, t):
    P(doc, t, size=16, bold=True, before=8, after=4)


def body(doc, t):
    P(doc, t, align="justify", first=1.0, after=6, spacing=1.5)


def br(doc):
    doc.add_page_break()


def set_row_height(row, cm):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn("w:val"), str(int(cm * 567)))
    trHeight.set(qn("w:hRule"), "atLeast")
    trPr.append(trHeight)


def table(doc, no, title, headers, rows):
    P(doc, f"ตารางที่ {no}  {title}", size=14, bold=True, align="center", before=6, after=4)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(h), 12, True)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.rows[r + 1].cells[c]
            cell.text = ""
            font(cell.paragraphs[0].add_run(str(val)), 12)
    P(doc, "", after=4)


def fig(doc, no, title, source, height=5.2, channel=""):
    P(doc, f"ภาพที่ {no}  {title}", size=13, bold=True, align="center", before=6, after=3)
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_row_height(t.rows[0], height)
    cell = t.rows[0].cells[0]
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run("ใส่ภาพตรงนี้"), 12)
    P(doc, f"แหล่งภาพ: {source}", size=11, align="center", after=2)
    if channel:
        P(doc, channel, size=12, align="center", after=6)


def toc_line(doc, t, *, indent=0, bold=False):
    P(doc, t, size=16, bold=bold, indent=indent, after=2, spacing=1.15)


def bib(doc, items):
    for r in items:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(1.27)
        pf.first_line_indent = Cm(-1.27)
        pf.space_after = Pt(10)
        pf.line_spacing = 1.5
        font(p.add_run(r), 16)


def build():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.54)
        s.bottom_margin = Cm(2.54)
        s.left_margin = Cm(3.0)
        s.right_margin = Cm(2.5)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(16)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    # หน้า 1 ปก
    seal = SEAL if os.path.exists(SEAL) else SEAL_FB
    if os.path.exists(seal):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        p.add_run().add_picture(seal, width=Cm(3.6))
    P(doc, "รายงานวิชาสื่อดิจิทัล", size=22, bold=True, align="center", after=8)
    P(doc, "การสร้างสรรค์สื่อเพื่อธุรกิจ Digital Gift", size=18, bold=True, align="center", after=4)
    P(doc, "โปสเตอร์บน Facebook และคลิปสั้นบน TikTok", size=16, align="center", after=16)
    P(doc, "รายวิชา สื่อดิจิทัล", size=15, align="center", after=4)
    P(doc, "ภาคเรียนที่ 1 ปีการศึกษา 2569", size=15, align="center", after=12)
    P(doc, "หลักสูตรเทคโนโลยีบัณฑิต  ภาควิชาเทคโนโลยีธุรกิจดิจิทัล", size=14, align="center", after=4)
    P(doc, "วิทยาลัยอาชีวศึกษาขอนแก่น", size=14, align="center", after=4)
    P(doc, "สถาบันการอาชีวศึกษาภาคตะวันออกเฉียงเหนือ 3", size=14, align="center", after=16)
    P(doc, f"จัดทำโดย  {STUDENT}", size=16, bold=True, align="center", after=4)
    P(doc, f"รหัสนักศึกษา  {SID}", size=14, align="center", after=10)
    P(doc, f"อาจารย์ประจำวิชา  {TEACHER}", size=14, align="center", after=4)

    # หน้า 2 คำนำ
    br(doc)
    H1(doc, "คำนำ")
    body(
        doc,
        "รายงานฉบับนี้จัดทำขึ้นเพื่อประกอบรายวิชาสื่อดิจิทัล ภาคเรียนที่ 1 ปีการศึกษา 2569 "
        "หลักสูตรเทคโนโลยีบัณฑิต ภาควิชาเทคโนโลยีธุรกิจดิจิทัล วิทยาลัยอาชีวศึกษาขอนแก่น "
        "โดยมีอาจารย์ประจำวิชา คือ นางชญานิตย์ ภาสว่าง",
    )
    body(
        doc,
        "เนื้อหานำเสนอหลักการคิดสื่อสร้างสรรค์ทางธุรกิจดิจิทัล องค์ประกอบและรูปแบบของสื่อ "
        "รวมถึงเทคนิคการสร้างสรรค์ข้อความ ภาพนิ่ง ภาพเคลื่อนไหว เสียง และวิดีโอ "
        "แล้วนำไปใช้กับธุรกิจจำลอง Digital Gift ซึ่งเป็นแพลตฟอร์มของขวัญดิจิทัลเฉพาะบุคคล",
    )
    body(
        doc,
        "ผลงานสื่อในรายงานนี้มี 2 ชิ้น คือ โปสเตอร์ลง Facebook และคลิปสั้นลง TikTok "
        "ทั้งสองชิ้นใช้แนวคิดเดียวกัน คือทำให้ของขวัญดิจิทัลรู้สึกพิเศษ เข้าใจง่าย และชวนไปดูตัวอย่าง",
    )
    body(
        doc,
        "ผู้จัดทำหวังว่ารายงานฉบับนี้จะเป็นประโยชน์ต่อการเรียนรู้ตามรายวิชา "
        "และขอขอบพระคุณอาจารย์ประจำวิชาที่ให้คำแนะนำจนรายงานสำเร็จลุล่วง",
    )
    P(doc, "", after=12)
    P(doc, STUDENT, align="right", after=2)
    P(doc, f"รหัสนักศึกษา {SID}", align="right", after=2)
    P(doc, "ภาคเรียนที่ 1 ปีการศึกษา 2569", align="right", after=2)

    # หน้า 3 สารบัญ
    br(doc)
    H1(doc, "สารบัญ")
    toc_line(doc, "ปก")
    toc_line(doc, "คำนำ")
    toc_line(doc, "สารบัญ")
    toc_line(doc, "สารบัญภาพ")
    toc_line(doc, "สารบัญตาราง")
    toc_line(doc, "บทที่ 1  หลักการคิดสื่อสร้างสรรค์ทางธุรกิจดิจิทัล", bold=True)
    toc_line(doc, "1.1  ความหมายและความสำคัญ", indent=0.75)
    toc_line(doc, "1.2  หลักการคิดสื่อสร้างสรรค์", indent=0.75)
    toc_line(doc, "บทที่ 2  องค์ประกอบและรูปแบบของสื่อดิจิทัล", bold=True)
    toc_line(doc, "2.1  องค์ประกอบของสื่อดิจิทัล", indent=0.75)
    toc_line(doc, "2.2  รูปแบบของสื่อดิจิทัล", indent=0.75)
    toc_line(doc, "บทที่ 3  เทคนิคการสร้างสรรค์สื่อ", bold=True)
    toc_line(doc, "3.1  เทคนิคการสร้างสรรค์ข้อความและภาพนิ่ง", indent=0.75)
    toc_line(doc, "3.2  เทคนิคการสร้างสรรค์ภาพเคลื่อนไหวและเสียง", indent=0.75)
    toc_line(doc, "3.3  เทคนิคการสร้างสรรค์วิดีโอ", indent=0.75)
    toc_line(doc, "สรุป", bold=True)
    toc_line(doc, "บรรณานุกรม", bold=True)

    P(doc, "สารบัญภาพ", size=16, bold=True, align="center", before=10, after=4)
    toc_line(doc, "ภาพที่ 1    โปสเตอร์ Facebook ของ Digital Gift")
    toc_line(doc, "ภาพที่ 2    คลิปสั้น TikTok ของ Digital Gift")

    P(doc, "สารบัญตาราง", size=16, bold=True, align="center", before=8, after=4)
    toc_line(doc, "ตารางที่ 1    องค์ประกอบสื่อที่ใช้ในงานนี้")
    toc_line(doc, "ตารางที่ 2    ชื่อช่องทางสื่อที่ใช้เผยแพร่")

    # หน้า 4 บทที่ 1
    br(doc)
    H1(doc, "บทที่ 1")
    P(doc, "หลักการคิดสื่อสร้างสรรค์ทางธุรกิจดิจิทัล", size=18, bold=True, align="center", after=8)

    H2(doc, "1.1  ความหมายและความสำคัญ")
    body(
        doc,
        "สื่อสร้างสรรค์ทางธุรกิจดิจิทัล คือ การออกแบบเนื้อหาเพื่อสื่อสารสินค้าผ่านช่องทางออนไลน์ "
        "ให้ผู้ชมเข้าใจเร็ว รู้สึกสนใจ และรู้ว่าต้องทำอะไรต่อ สื่อที่ดีจึงมีแนวคิดชัด ไม่ได้ทำเพียงให้สวย",
    )
    body(
        doc,
        "สื่อดิจิทัลมีความสำคัญเพราะช่วยให้ธุรกิจเข้าถึงลูกค้าได้ตรงกลุ่ม ใช้ต้นทุนต่ำ และส่งต่อได้ง่าย "
        "สำหรับ Digital Gift สื่อทำให้คนเห็นตัวอย่างของขวัญดิจิทัลก่อนตัดสินใจทัก LINE",
    )

    H2(doc, "1.2  หลักการคิดสื่อสร้างสรรค์")
    body(
        doc,
        "งานนี้ใช้หลักสามข้อ คือ เปิดให้สะดุดตา อธิบายคุณค่าของสินค้าด้วยข้อความสั้น "
        "และปิดด้วยคำชวนไปดูตัวอย่าง ทั้งโปสเตอร์ Facebook และคลิปสั้น TikTok ใช้หลักเดียวกัน",
    )
    body(
        doc,
        "แนวคิดหลักคือ “ของขวัญดิจิทัลที่รู้สึกพิเศษจริง ๆ” "
        "กลุ่มเป้าหมายคือวัยรุ่นถึงวัยทำงานที่อยากเซอร์ไพรส์คนสำคัญ "
        "สิ่งที่อยากให้ผู้ชมทำต่อคือเปิดดูตัวอย่างบนเว็บ หรือทัก LINE เพื่อสั่งทำ",
    )

    # หน้า 5-6 บทที่ 2
    br(doc)
    H1(doc, "บทที่ 2")
    P(doc, "องค์ประกอบและรูปแบบของสื่อดิจิทัล", size=18, bold=True, align="center", after=8)

    H2(doc, "2.1  องค์ประกอบของสื่อดิจิทัล")
    body(
        doc,
        "สื่อดิจิทัลมี 5 องค์ประกอบ คือ ข้อความ ภาพนิ่ง ภาพเคลื่อนไหว เสียง และวิดีโอ "
        "งานนี้เลือกใช้ตามช่องทาง โปสเตอร์เน้นข้อความกับภาพนิ่ง คลิปสั้นใช้ครบทุกองค์ประกอบ",
    )
    table(
        doc,
        1,
        "องค์ประกอบสื่อที่ใช้ในงานนี้",
        ["องค์ประกอบ", "โปสเตอร์ Facebook", "คลิปสั้น TikTok"],
        [
            ["ข้อความ", "ใช้", "ใช้"],
            ["ภาพนิ่ง", "ใช้", "ใช้"],
            ["ภาพเคลื่อนไหว", "ไม่ใช้", "ใช้"],
            ["เสียง", "ไม่ใช้", "ใช้"],
            ["วิดีโอ", "ไม่ใช้", "ใช้"],
        ],
    )
    body(
        doc,
        "ข้อความใช้บอกความหมายหลัก ภาพนิ่งช่วยให้จำสินค้าได้ "
        "ภาพเคลื่อนไหวและเสียงใช้ดึงความสนใจในคลิปสั้น ส่วนวิดีโอใช้เล่าเรื่องให้จบในเวลาสั้น",
    )

    br(doc)
    H2(doc, "2.2  รูปแบบของสื่อดิจิทัล")
    body(
        doc,
        "งานนี้ใช้สื่อ 2 รูปแบบ คือโปสเตอร์ภาพนิ่งบน Facebook และคลิปสั้นแนวตั้งบน TikTok "
        "Facebook เหมาะกับการอ่านและแชร์ TikTok เหมาะกับการหยุดดูคลิปสั้น "
        "ชื่อเพจ Facebook และชื่อบัญชี TikTok ให้เขียนลงในตารางที่ 2 และบรรทัดใต้ภาพประกอบ",
    )
    table(
        doc,
        2,
        "ชื่อช่องทางสื่อที่ใช้เผยแพร่",
        ["ช่องทาง", "ชื่อช่องที่ใช้จริง", "ใช้เผยแพร่"],
        [
            ["Facebook", "เขียนชื่อเพจตรงนี้ ........................", "โปสเตอร์"],
            ["TikTok", "เขียนชื่อบัญชีตรงนี้ ......................", "คลิปสั้น"],
        ],
    )
    fig(
        doc,
        1,
        "โปสเตอร์ Facebook ของ Digital Gift",
        "ผู้จัดทำ / Facebook",
        height=4.2,
        channel="ชื่อเพจ Facebook: ................................................",
    )
    fig(
        doc,
        2,
        "คลิปสั้น TikTok ของ Digital Gift",
        "TikTok",
        height=4.2,
        channel="ชื่อบัญชี TikTok: ................................................",
    )

    # หน้า 7-9 บทที่ 3
    br(doc)
    H1(doc, "บทที่ 3")
    P(doc, "เทคนิคการสร้างสรรค์สื่อ", size=18, bold=True, align="center", after=8)

    H2(doc, "3.1  เทคนิคการสร้างสรรค์ข้อความและภาพนิ่ง")
    body(
        doc,
        "ข้อความเขียนสั้น ชัด และพูดแบบคุย ไม่ใช้คำโฆษณาที่ยาก "
        "หัวข้อหลักคือ “ของขวัญดิจิทัลที่รู้สึกพิเศษจริง ๆ” "
        "ข้อความรองคือ “เว็บไซต์เฉพาะบุคคล ส่งผ่านลิงก์เดียว” "
        "คำชวนคือ “ดูตัวอย่างได้เลย”",
    )
    body(
        doc,
        "ภาพนิ่งจัดให้มีจุดโฟกัสเดียว ตัวอักษรตัดกับพื้นหลังให้อ่านง่ายบนมือถือ "
        "เลือกโทนอบอุ่นให้เข้ากับสินค้าที่เป็นของขวัญ และไม่ใส่รายละเอียดจนรก "
        "โปสเตอร์ใช้ภาพนิ่งเป็นชิ้นงานหลัก จากนั้นนำไปตัดต่อในคลิปสั้นด้วย",
    )

    br(doc)
    H2(doc, "3.2  เทคนิคการสร้างสรรค์ภาพเคลื่อนไหวและเสียง")
    body(
        doc,
        "ภาพเคลื่อนไหวใช้เฉพาะคลิป TikTok เพื่อดึงสายตาในช่วงวินาทีแรก "
        "เทคนิคที่ใช้ได้แก่ การซูมเข้าภาพของขวัญ การเลื่อนข้อความ และการตัดต่อฉากให้สั้น "
        "ไม่ใส่เอฟเฟกต์มากจนแย่งความสนใจจากสินค้า",
    )
    body(
        doc,
        "เสียงใช้เพลงเบา ๆ ให้เข้ากับความรู้สึกอบอุ่น และไม่ดังรบกวนข้อความ "
        "โปสเตอร์ไม่มีเสียง จึงต้องสื่อสารให้จบด้วยข้อความและภาพนิ่ง "
        "คลิปสั้นมีข้อความบนจอเสมอ เพราะผู้ชมอาจดูแบบปิดเสียง",
    )

    br(doc)
    H2(doc, "3.3  เทคนิคการสร้างสรรค์วิดีโอ")
    body(
        doc,
        "คลิปสั้นทำแนวตั้ง ความยาวไม่มาก ดูจบในรอบเดียว "
        "โครงคลิปมีสามช่วง คือเปิดให้สะดุดตา โชว์ตัวอย่างสินค้า และปิดด้วยคำชวนไปดูเดโม",
    )
    body(
        doc,
        "ช่วงเปิดใช้ภาพหรือคำถามให้คนหยุดดู ช่วงกลางโชว์หน้าตาของขวัญดิจิทัล "
        "ช่วงปิดบอกว่าสั่งทำง่าย เพราะ Digital Gift เป็นงานสั่งทำรายชิ้น ไม่มีตะกร้าสินค้า",
    )
    body(
        doc,
        "สรุปเทคนิคทั้งบทนี้คือ ข้อความสั้น ภาพไม่รก จังหวะเปิดคลิปชัด และมีคำชวนที่ทำตามได้ "
        "โปสเตอร์ให้ข้อมูลชัด คลิปสั้นดึงความสนใจ ทั้งสองชิ้นสื่อสารแนวคิดเดียวกัน",
    )

    # หน้า 10 สรุป
    br(doc)
    H1(doc, "สรุป")
    body(
        doc,
        "รายงานฉบับนี้นำหลักการคิดสื่อสร้างสรรค์ทางธุรกิจดิจิทัลมาใช้กับ Digital Gift "
        "โดยกำหนดแนวคิดให้ชัดว่าของขวัญดิจิทัลต้องรู้สึกพิเศษ เข้าใจง่าย และส่งต่อได้ด้วยลิงก์",
    )
    body(
        doc,
        "องค์ประกอบสื่อถูกเลือกตามช่องทาง โปสเตอร์ Facebook ใช้ข้อความกับภาพนิ่งเป็นหลัก "
        "คลิปสั้น TikTok ใช้ภาพเคลื่อนไหว เสียง และวิดีโอเป็นหลัก "
        "ทั้งสองชิ้นมีรูปแบบต่างกัน แต่สื่อสารแนวคิดเดียวกัน",
    )
    body(
        doc,
        "เทคนิคที่ใช้จึงไม่ซับซ้อน คือเขียนข้อความสั้น จัดภาพให้มีจุดโฟกัสเดียว "
        "ใช้ภาพเคลื่อนไหวและเสียงเท่าที่จำเป็น และทำคลิปสั้นให้มีคำชวนไปดูตัวอย่าง "
        "สื่อทั้งสองชิ้นจึงเหมาะกับการนำเสนอธุรกิจจำลองในรายวิชาสื่อดิจิทัล",
    )
    body(
        doc,
        "ผู้จัดทำสรุปได้ว่า การสร้างสรรค์สื่อทางธุรกิจดิจิทัลต้องเริ่มจากแนวคิดและกลุ่มเป้าหมาย "
        "แล้วจึงเลือกองค์ประกอบกับช่องทางให้เหมาะกับสินค้า ไม่ใช่ทำสื่อหลายชิ้นโดยไม่มีทิศทางเดียวกัน",
    )

    # หน้า 11 บรรณานุกรม
    br(doc)
    H1(doc, "บรรณานุกรม")
    P(doc, "", after=8)
    bib(
        doc,
        [
            "Chaffey, D., & Ellis-Chadwick, F. (2022). Digital marketing (8th ed.). Pearson.",
            "Kotler, P., Kartajaya, H., & Setiawan, I. (2021). Marketing 5.0: Technology for humanity. Wiley.",
            "Meta. (2024). Facebook creative best practices. Meta Business Help Center. Retrieved September 21, 2026, from https://www.facebook.com/business/help",
            "TikTok. (2024). Creative best practices for short-form video. TikTok Creative Center. Retrieved September 21, 2026, from https://ads.tiktok.com/help/",
            "สำนักงานพัฒนาธุรกรรมทางอิเล็กทรอนิกส์. (2566). รายงานพฤติกรรมผู้ใช้อินเทอร์เน็ตในประเทศไทย ปี 2566. สพธอ. สืบค้น 21 กันยายน 2569, จาก https://www.etda.or.th",
        ],
    )

    saved = []
    for path in OUTS:
        try:
            os.makedirs(os.path.dirname(path), exist_ok=True)
            doc.save(path)
            saved.append(path)
        except OSError as e:
            print(f"skip {path}: {e}")
    return saved


if __name__ == "__main__":
    for p in build():
        print("saved", p)
