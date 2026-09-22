# -*- coding: utf-8 -*-
"""เล่มรายงาน Digital Gift แบบกระชับ — ยกตัวอย่าง Demo เดียว"""
from __future__ import annotations

import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = os.path.dirname(__file__)
OUTS = [
    os.path.join(BASE, "รายงานโครงการ-Digital-Gift-กระชับ.docx"),
    os.path.join(os.path.expanduser("~"), "Downloads", "รายงานโครงการ-Digital-Gift-กระชับ.docx"),
]
QR = os.path.join(BASE, "assets", "digital-gift-qr.png")
SEAL = os.path.join(BASE, "assets", "ivec-seal-user.png")
SEAL_FB = os.path.join(BASE, "assets", "ivec3-seal.png")
WEB = "https://parapon12.github.io/Digital_Gift/"
REPO = "https://github.com/Parapon12/Digital_Gift"
LINE = "https://lin.ee/uoPtq9r"
FONT = "TH Sarabun New"
STUDENT = "นายพลพล กำมะถัน"
SID = "69419100009"
TEACHER = "นางชญานิตย์ ภาสว่าง"


def font(run, size=16, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), FONT)


def P(doc, text="", *, size=16, bold=False, align="left", before=0, after=6, first=0, indent=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first:
        pf.first_line_indent = Cm(first)
    if indent:
        pf.left_indent = Cm(indent)
    p.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    if text:
        font(p.add_run(text), size=size, bold=bold)
    return p


def H1(doc, t):
    P(doc, t, size=20, bold=True, align="center", before=14, after=10)


def H2(doc, t):
    P(doc, t, size=17, bold=True, before=10, after=6)


def body(doc, t):
    P(doc, t, align="justify", first=1.0, after=6)


def num(doc, items):
    for i, t in enumerate(items, 1):
        P(doc, f"{i}.  {t}", align="justify", indent=0.75, after=3)


def dot(doc, items):
    for t in items:
        P(doc, f"•  {t}", align="justify", indent=0.75, after=3)


def br(doc):
    doc.add_page_break()


def table(doc, no, title, headers, rows, note=""):
    P(doc, f"ตารางที่ {no}  {title}", size=15, bold=True, align="center", before=8, after=4)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(h), 13, True)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.rows[r + 1].cells[c]
            cell.text = ""
            font(cell.paragraphs[0].add_run(str(val)), 12)
    P(doc, "", after=2)
    if note:
        P(doc, note, size=12, after=6)


def fig(doc, no, title, source):
    P(doc, f"[ แทรกภาพที่ {no} ]", size=13, bold=True, align="center", before=6, after=2)
    P(doc, f"ภาพที่ {no}  {title}", size=13, bold=True, align="center", after=1)
    P(doc, f"แหล่งภาพ: {source}", size=11, align="center", after=8)


def build():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.54)
        s.bottom_margin = Cm(2.54)
        s.left_margin = Cm(3.0)
        s.right_margin = Cm(2.5)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(16)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    # ----- ปกหน้า -----
    H1(doc, "รายงานโครงการ")
    P(doc, "การพัฒนาเว็บอีคอมเมิร์ซ", size=18, bold=True, align="center", after=4)
    P(doc, "โครงการ Digital Gift", size=20, bold=True, align="center", after=4)
    P(doc, "แพลตฟอร์มของขวัญดิจิทัลเฉพาะบุคคล", size=16, align="center", after=14)
    P(doc, "รายวิชา การพัฒนาพาณิชย์อิเล็กทรอนิกส์สำหรับธุรกิจดิจิทัล", size=14, align="center", after=2)
    P(doc, "รหัสรายวิชา 26-41910-2101  |  ภาคเรียนที่ 1 ปีการศึกษา 2569", size=14, align="center", after=10)
    P(doc, "หลักสูตรเทคโนโลยีบัณฑิต  ภาควิชาเทคโนโลยีธุรกิจดิจิทัล", size=14, align="center", after=2)
    P(doc, "วิทยาลัยอาชีวศึกษาขอนแก่น", size=14, align="center", after=2)
    P(doc, "สถาบันการอาชีวศึกษาภาคตะวันออกเฉียงเหนือ 3", size=14, align="center", after=12)
    P(doc, f"จัดทำโดย  {STUDENT}", size=15, bold=True, align="center", after=2)
    P(doc, f"รหัสนักศึกษา  {SID}", size=14, align="center", after=8)
    P(doc, f"อาจารย์ประจำวิชา  {TEACHER}", size=14, align="center", after=4)

    # ----- รองปกใน -----
    br(doc)
    seal = SEAL if os.path.exists(SEAL) else SEAL_FB
    if os.path.exists(seal):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(seal, width=Cm(4.2))
        P(doc, "", after=6)
    P(doc, "สถาบันการอาชีวศึกษาภาคตะวันออกเฉียงเหนือ 3", size=15, bold=True, align="center", after=2)
    P(doc, "วิทยาลัยอาชีวศึกษาขอนแก่น", size=14, align="center", after=12)
    P(doc, "รายงานโครงการ Digital Gift", size=18, bold=True, align="center", after=6)
    P(doc, f"ผู้จัดทำ  {STUDENT}   รหัส  {SID}", size=14, align="center", after=4)
    P(doc, f"อาจารย์ประจำวิชา  {TEACHER}", size=14, align="center", after=4)
    P(doc, "ภาคเรียนที่ 1 ปีการศึกษา 2569", size=14, align="center", after=4)

    # ----- คำนำ -----
    br(doc)
    H1(doc, "คำนำ")
    body(
        doc,
        "รายงานฉบับนี้จัดทำเพื่อประกอบรายวิชา การพัฒนาพาณิชย์อิเล็กทรอนิกส์สำหรับธุรกิจดิจิทัล "
        "ภาคเรียนที่ 1 ปีการศึกษา 2569 โดยนำเสนอต้นแบบเว็บไซต์ Digital Gift ซึ่งเป็นแพลตฟอร์มของขวัญดิจิทัลเฉพาะบุคคล "
        "รูปแบบธุรกิจเป็น B2C ลูกค้าทดลองเล่น Demo บนเว็บ สั่งทำผ่าน LINE แล้วรับของขวัญเป็นลิงก์",
    )
    body(
        doc,
        "เอกสารจัดทำแบบกระชับ เน้นหลักการ ขั้นตอนพัฒนา และยกตัวอย่างการใช้งานซื้อขายเพียง 1 เทมเพลต "
        "คือหน้ารำลึกความทรงจำ เพื่อให้เห็นภาพการใช้งานจริงโดยไม่ซ้ำซ้อน",
    )
    P(doc, "", after=10)
    P(doc, STUDENT, align="right", after=1)
    P(doc, f"รหัสนักศึกษา {SID}", align="right", after=1)
    P(doc, "ภาคเรียนที่ 1 ปีการศึกษา 2569", align="right", after=2)

    # ----- สารบัญ -----
    br(doc)
    H1(doc, "สารบัญ")
    for t in [
        "คำนำ",
        "สารบัญ",
        "สารบัญภาพ",
        "สารบัญตาราง",
        "1.  หลักการและเหตุผล",
        "2.  ความหมายของธุรกิจอีคอมเมิร์ซ",
        "3.  ความแตกต่างระหว่าง E-Business และ E-Commerce",
        "4.  ประเภทของธุรกิจอีคอมเมิร์ซ",
        "5.  เครื่องมือและขั้นตอนการพัฒนาเว็บ Digital Gift",
        "6.  สรุปผล",
        "บรรณานุกรม",
        "ภาคผนวก  คู่มือและตัวอย่างการใช้งาน (หน้ารำลึกความทรงจำ)",
        "ภาคผนวก  ลิงก์เว็บไซต์และ QR Code",
    ]:
        P(doc, t, size=15, after=3)

    br(doc)
    H1(doc, "สารบัญภาพ")
    for no, title in [
        (1, "หน้าแรกเว็บไซต์ Digital Gift"),
        (2, "การ์ดเทมเพลตหน้ารำลึกความทรงจำ"),
        (3, "หน้า Demo หน้ารำลึกความทรงจำ"),
        (4, "ตัวอย่างช่วงความทรงจำใน Demo"),
        (5, "การเปิดข้อความลับใน Demo"),
        (6, "การสั่งทำผ่าน LINE"),
        (7, "หน้าแอดมินกรอกเนื้อหา"),
        (8, "การส่งมอบด้วยลิงก์"),
        (9, "ผู้รับเปิดลิงก์ของขวัญบนมือถือ"),
        (10, "QR Code เว็บไซต์สาธิต"),
    ]:
        P(doc, f"ภาพที่ {no}    {title}", size=14, after=2)

    br(doc)
    H1(doc, "สารบัญตาราง")
    for no, title in [
        (1, "ความแตกต่างระหว่าง E-Commerce และ E-Business"),
        (2, "ประเภทธุรกิจอีคอมเมิร์ซและความเกี่ยวข้องกับ Digital Gift"),
        (3, "เทคโนโลยีที่ใช้พัฒนา Digital Gift"),
        (4, "ขั้นตอนลูกค้าตั้งแต่เข้าเว็บจนได้ลิงก์ (ตัวอย่าง 1 เทมเพลต)"),
    ]:
        P(doc, f"ตารางที่ {no}    {title}", size=14, after=2)

    # ----- 1 -----
    br(doc)
    H1(doc, "1.  หลักการและเหตุผล")
    body(
        doc,
        "ของขวัญในรูปแบบข้อความแชทหรือรูปภาพเดี่ยวมักจางหายได้ง่าย ผู้จัดทำจึงพัฒนาต้นแบบ Digital Gift "
        "ให้ลูกค้าทดลองเล่นตัวอย่างบนเว็บ แล้วสั่งทำของขวัญดิจิทัลเฉพาะบุคคลผ่าน LINE "
        "ร้านจัดทำในระบบแอดมินและส่งมอบด้วยลิงก์ โดยยังเป็นต้นแบบเพื่อการเรียนรู้และสาธิต ยังไม่ใช่เว็บบริการเชิงพาณิชย์จริง",
    )
    H2(doc, "1.1  วัตถุประสงค์")
    num(
        doc,
        [
            "ศึกษาความหมาย ประเภท และองค์ประกอบของธุรกิจอีคอมเมิร์ซ",
            "พัฒนาต้นแบบ Digital Gift ด้วยแนวทาง Custom Development",
            "ออกแบบกระบวนการลูกค้าตั้งแต่ Demo จนถึงส่งมอบด้วยลิงก์",
        ],
    )
    H2(doc, "1.2  ลิงก์เว็บไซต์สาธิต")
    P(doc, f"เว็บไซต์:  {WEB}", after=2)
    P(doc, f"คลังโค้ด:  {REPO}", after=2)
    P(doc, f"LINE สั่งทำ:  {LINE}", after=4)
    P(doc, "หมายเหตุ: นำลิงก์เว็บไปวางบน Padlet หรือกระดานออนไลน์ที่ผู้สอนกำหนด", size=13, after=6)

    # ----- 2 -----
    br(doc)
    H1(doc, "2.  ความหมายของธุรกิจอีคอมเมิร์ซ")
    body(
        doc,
        "ธุรกิจอีคอมเมิร์ซ (E-Commerce) คือการซื้อขายสินค้าหรือบริการผ่านเครือข่ายอิเล็กทรอนิกส์ "
        "ครอบคลุมการนำเสนอสินค้า การรับคำสั่งซื้อ การชำระเงิน และการส่งมอบ ในโครงงาน Digital Gift "
        "สินค้าคือของขวัญดิจิทัลตามออเดอร์ และส่งมอบผ่านลิงก์เว็บไซต์เฉพาะบุคคล",
    )

    # ----- 3 -----
    br(doc)
    H1(doc, "3.  ความแตกต่างระหว่าง E-Business และ E-Commerce")
    table(
        doc,
        1,
        "ความแตกต่างระหว่าง E-Commerce และ E-Business",
        ["หัวข้อ", "E-Commerce", "E-Business"],
        [
            ["ความหมาย", "เน้นการซื้อขายออนไลน์", "การใช้ดิจิทัลในกระบวนการธุรกิจโดยรวม"],
            ["ขอบเขต", "แคบกว่า", "กว้างกว่า"],
            ["Digital Gift", "อยู่ในกรอบนี้เป็นหลัก", "สัมพันธ์ระดับต้นแบบผ่านระบบแอดมิน"],
        ],
    )

    # ----- 4 -----
    br(doc)
    H1(doc, "4.  ประเภทของธุรกิจอีคอมเมิร์ซ")
    body(doc, "จำแนกตามคู่ผู้ซื้อ–ผู้ขายได้ 4 ประเภทหลัก Digital Gift ออกแบบเป็น B2C")
    table(
        doc,
        2,
        "ประเภทธุรกิจอีคอมเมิร์ซและความเกี่ยวข้องกับ Digital Gift",
        ["ประเภท", "คู่ธุรกรรม", "เกี่ยวข้องกับ Digital Gift"],
        [
            ["B2B", "ธุรกิจ → ธุรกิจ", "ไม่ใช่รูปแบบหลัก"],
            ["B2C", "ธุรกิจ → ผู้บริโภค", "ใช่ — รูปแบบหลัก"],
            ["C2C", "ผู้บริโภค → ผู้บริโภค", "ไม่ใช่รูปแบบหลัก"],
            ["C2B", "ผู้บริโภค → ธุรกิจ", "ไม่ใช่รูปแบบหลัก"],
        ],
    )

    # ----- 5 -----
    br(doc)
    H1(doc, "5.  เครื่องมือและขั้นตอนการพัฒนาเว็บ Digital Gift")
    body(doc, "ผู้จัดทำเลือกแนวทาง Custom Development เพื่อควบคุมประสบการณ์ของขวัญแบบอินเทอร์แอคทีฟได้เอง")
    table(
        doc,
        3,
        "เทคโนโลยีที่ใช้พัฒนา Digital Gift",
        ["ชั้นระบบ", "เทคโนโลยี", "บทบาท"],
        [
            ["Frontend", "React + TypeScript + Vite", "หน้าเว็บ Demo และของขวัญ"],
            ["Backend", "Go + Chi + JWT", "API และระบบแอดมิน"],
            ["Database", "PostgreSQL 16", "เก็บของขวัญและเดโม"],
            ["ช่องทางสั่งทำ", "LINE (ลิงก์ OA)", "รับออเดอร์และสื่อสาร"],
        ],
    )

    H2(doc, "5.1  ขั้นตอนการพัฒนา")
    num(
        doc,
        [
            "ศึกษาความต้องการและกำหนดแนวคิด B2C ของขวัญดิจิทัล",
            "ออกแบบ Customer Journey: Demo → LINE → จัดทำ → จ่ายเงิน → ส่งลิงก์",
            "พัฒนา Frontend / Backend / ฐานข้อมูล",
            "ทดสอบ Demo และการเปิดลิงก์ของขวัญบนมือถือ",
            f"เผยแพร่หน้าสาธิตที่ {WEB} และจัดทำรายงาน",
        ],
    )
    H2(doc, "5.2  สรุปกระบวนการลูกค้า")
    body(
        doc,
        "เข้าเว็บ → เล่น Demo → ทัก LINE สั่งทำ → ส่งชื่อ ข้อความ รูป → ร้านจัดทำ → ส่ง QR ชำระเงิน → "
        "ลูกค้าโอนแล้ว → ร้านส่งลิงก์ของขวัญให้มอบผู้รับ",
    )

    # ----- 6 -----
    br(doc)
    H1(doc, "6.  สรุปผล")
    body(
        doc,
        "โครงงาน Digital Gift สำเร็จในฐานะต้นแบบเว็บอีคอมเมิร์ซของขวัญดิจิทัลเฉพาะบุคคล "
        "สามารถอธิบายหลักการอีคอมเมิร์ซ พัฒนาด้วย Custom Development และสาธิตกระบวนการสั่งทำจนส่งมอบด้วยลิงก์ได้",
    )

    # ----- บรรณานุกรม -----
    br(doc)
    H1(doc, "บรรณานุกรม")
    bibs = [
        "Bendersky, E., & The Go Team. (2024, February 6). Go 1.22 is released! The Go Project. Retrieved September 19, 2026, from https://go.dev/blog/go1.22",
        "Jones, M., Bradley, J., & Sakimura, N. (2015, May). JSON Web Token (JWT) (RFC 7519). Internet Engineering Task Force. Retrieved September 19, 2026, from https://www.rfc-editor.org/rfc/rfc7519",
        "PostgreSQL Global Development Group. (2023, September 14). PostgreSQL 16 Released! PostgreSQL. Retrieved September 19, 2026, from https://www.postgresql.org/about/news/postgresql-16-released-2715/",
        "The React Team. (2024, December 5). React v19. Meta Platforms. Retrieved September 19, 2026, from https://react.dev/blog/2024/12/05/react-19",
        "The Vite Team. (2026, March 12). Vite 8.0 is out! Vite.js. Retrieved September 19, 2026, from https://vite.dev/blog/announcing-vite8",
    ]
    for r in bibs:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(1.27)
        pf.first_line_indent = Cm(-1.27)
        pf.space_after = Pt(8)
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        font(p.add_run(r), 15)

    # ----- ภาคผนวก: ตัวอย่างเดียว -----
    br(doc)
    H1(doc, "ภาคผนวก")
    P(doc, "คู่มือและตัวอย่างการใช้งานซื้อขายผ่านเว็บ", size=16, bold=True, align="center", after=8)
    body(
        doc,
        "ภาคผนวกนี้ยกตัวอย่างเพียง 1 เทมเพลต คือ หน้ารำลึกความทรงจำ "
        "ราคาจำลอง 79 บาท เส้นทาง Demo: /demo/memory-page เพื่อแสดงขั้นตอนใช้งานจริงตั้งแต่เล่นตัวอย่างจนได้รับลิงก์",
    )

    H2(doc, "ก.1  ตัวอย่างที่เลือก")
    dot(
        doc,
        [
            "ชื่อเทมเพลต: หน้ารำลึกความทรงจำ",
            "ราคาจำลอง: 79 บาท",
            f"Demo: {WEB}demo/memory-page",
            "ลักษณะ: scrapbook เลื่อนดูรูป แคปชัน และเปิดข้อความลับ",
        ],
    )

    H2(doc, "ก.2  ขั้นตอนทดลองเล่น Demo")
    num(
        doc,
        [
            "เข้าหน้าแรกของเว็บไซต์",
            "เลือกการ์ด “หน้ารำลึกความทรงจำ”",
            "เปิดหน้า Demo แล้วเลื่อนดูช่วงความทรงจำ",
            "กดเปิดซองหรือข้อความลับเพื่ออ่านเนื้อหา",
        ],
    )
    fig(doc, 1, "หน้าแรกเว็บไซต์ Digital Gift", "หน้า /")
    fig(doc, 2, "การ์ดเทมเพลตหน้ารำลึกความทรงจำ", "หน้า / ส่วนตัวอย่าง")
    fig(doc, 3, "หน้า Demo หน้ารำลึกความทรงจำ", "/demo/memory-page")
    fig(doc, 4, "ตัวอย่างช่วงความทรงจำใน Demo", "/demo/memory-page")
    fig(doc, 5, "การเปิดข้อความลับใน Demo", "/demo/memory-page")

    H2(doc, "ก.3  ขั้นตอนสั่งทำจนได้ลิงก์")
    table(
        doc,
        4,
        "ขั้นตอนลูกค้าตั้งแต่เข้าเว็บจนได้ลิงก์ (ตัวอย่าง 1 เทมเพลต)",
        ["ขั้น", "ผู้ทำ", "สิ่งที่เกิดขึ้น"],
        [
            ["1", "ลูกค้า", "เข้าเว็บและเล่น Demo หน้ารำลึกความทรงจำ"],
            ["2", "ลูกค้า", "ทัก LINE แจ้งเลือกแบบนี้ (79 บาท)"],
            ["3", "ร้าน", "ส่งรายการข้อมูลที่ต้องเตรียม"],
            ["4", "ลูกค้า", "ส่งชื่อ ผู้รับ/ผู้ส่ง ข้อความ และรูป → พิมพ์「ส่งครบแล้ว」"],
            ["5", "แอดมิน", "จัดทำของขวัญในระบบหลังบ้าน"],
            ["6", "ร้าน", "ส่ง QR ให้ชำระเงิน"],
            ["7", "ลูกค้า", "โอนเงิน พิมพ์「โอนแล้ว」+ ส่งสลิป"],
            ["8", "ร้าน", "ตรวจสลิปแล้วส่งลิงก์ /gift/... ให้ลูกค้า"],
        ],
        note="ลูกค้าได้รับของขวัญเมื่อเปิดลิงก์ได้สำเร็จ",
    )
    fig(doc, 6, "การสั่งทำผ่าน LINE", "แอป LINE")
    fig(doc, 7, "หน้าแอดมินกรอกเนื้อหา", "/admin")
    fig(doc, 8, "การส่งมอบด้วยลิงก์", "แอป LINE")
    fig(doc, 9, "ผู้รับเปิดลิงก์ของขวัญบนมือถือ", "เปิด /gift/... ")

    # ----- ลิงก์ + QR -----
    br(doc)
    H1(doc, "ภาคผนวก")
    P(doc, "ลิงก์เว็บไซต์และ QR Code", size=16, bold=True, align="center", after=8)
    P(doc, "เว็บไซต์สาธิต", size=14, bold=True, after=2)
    P(doc, WEB, after=6)
    P(doc, "คลังโค้ด", size=14, bold=True, after=2)
    P(doc, REPO, after=6)
    P(doc, "LINE สั่งทำ", size=14, bold=True, after=2)
    P(doc, LINE, after=8)
    P(doc, "ภาพที่ 10  QR Code สำหรับสแกนเข้าเว็บไซต์สาธิต", size=13, bold=True, align="center", after=6)
    if os.path.exists(QR):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(QR, width=Cm(4.5))
    P(doc, "ให้นำลิงก์เว็บไปวางบน Padlet ตามที่ผู้สอนกำหนด", size=12, align="center", after=8)

    # ----- ปกหลัง -----
    br(doc)
    for _ in range(5):
        P(doc, "", after=8)
    P(doc, "Digital Gift", size=20, bold=True, align="center", after=6)
    P(doc, "แพลตฟอร์มของขวัญดิจิทัลเฉพาะบุคคล", size=14, align="center", after=10)
    P(doc, WEB, size=12, align="center", after=8)
    P(doc, "วิทยาลัยอาชีวศึกษาขอนแก่น", size=13, align="center", after=2)
    P(doc, "สถาบันการอาชีวศึกษาภาคตะวันออกเฉียงเหนือ 3", size=13, align="center", after=2)
    P(doc, "ภาคเรียนที่ 1 ปีการศึกษา 2569", size=13, align="center", after=6)
    P(doc, "เย็บเล่มด้วยสันพลาสติกหรือสันรูด", size=11, align="center", after=4)

    for path in OUTS:
        try:
            doc.save(path)
            print("SAVED", path)
        except PermissionError:
            alt = path.replace(".docx", "-ใหม่.docx")
            doc.save(alt)
            print("LOCKED->", alt)


if __name__ == "__main__":
    build()
