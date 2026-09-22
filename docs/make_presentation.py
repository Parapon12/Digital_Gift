# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
ACCENT = RGBColor(0xC4, 0x5C, 0x26)
CREAM = RGBColor(0xF7, 0xF3, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x2C, 0x2C, 0x2C)
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
LIGHT = RGBColor(0xEE, 0xE8, 0xDF)
SOFT = RGBColor(0xC8, 0xD0, 0xE0)


def set_run(run, size=18, bold=False, color=DARK, font="TH Sarabun New"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", font)


def add_bg(slide, color=CREAM):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_top_bar(slide):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.15))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def add_footer(slide, num, total=12):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(10), Inches(0.35))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f"Digital Gift  |  ต้นแบบแพลตฟอร์มของขวัญดิจิทัลเฉพาะบุคคล  |  {num}/{total}"
    set_run(run, 12, False, GRAY)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(6.95), Inches(1.2), Inches(0.04)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def title_box(slide, text, top=0.35):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(12), Inches(0.7))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run(run, 32, True, NAVY)


def body_box(slide, left, top, width, height):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def add_para(tf, text, size=20, bold=False, color=DARK, align=PP_ALIGN.LEFT, space_before=4):
    if tf.paragraphs[0].text == "" and len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)


def card(slide, left, top, width, height, fill=WHITE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = LIGHT
    shape.adjustments[0] = 0.08
    return shape


# ========== SLIDE 1 COVER ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
strip = slide.shapes.add_shape(
    MSO_SHAPE.RECTANGLE, 0, Inches(5.9), prs.slide_width, Inches(1.6)
)
strip.fill.solid()
strip.fill.fore_color.rgb = RGBColor(0x14, 0x20, 0x38)
strip.line.fill.background()
accent = slide.shapes.add_shape(
    MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(1.5), Inches(0.08)
)
accent.fill.solid()
accent.fill.fore_color.rgb = ACCENT
accent.line.fill.background()

tf = body_box(slide, 0.8, 1.7, 11.5, 4)
add_para(
    tf,
    "รายงานการพัฒนาพาณิชย์อิเล็กทรอนิกส์สำหรับธุรกิจดิจิทัล",
    18,
    False,
    SOFT,
    space_before=0,
)
add_para(tf, "Digital Gift", 48, True, WHITE, space_before=12)
add_para(tf, "ต้นแบบแพลตฟอร์มของขวัญดิจิทัลเฉพาะบุคคล", 24, False, RGBColor(0xE8, 0xD5, 0xC4), space_before=8)
add_para(
    tf,
    "Prototype / แบบจำลองระบบ  •  ยังไม่ใช่เว็บบริการเชิงพาณิชย์จริง",
    16,
    False,
    RGBColor(0xA8, 0xB4, 0xC8),
    space_before=16,
)

tf2 = body_box(slide, 0.8, 6.1, 11.5, 1.2)
add_para(
    tf2,
    "วิทยาลัยอาชีวศึกษาขอนแก่น  |  ภาคเรียนที่ 1 ปีการศึกษา 2569",
    16,
    False,
    SOFT,
    space_before=0,
)
add_para(tf2, "อาจารย์ที่ปรึกษา: นางชญานิตย์ ภาสว่าง", 16, False, SOFT, space_before=4)
add_para(
    tf2,
    "จัดทำโดย: ..............................    รหัสนักศึกษา: ..............................",
    16,
    False,
    WHITE,
    space_before=4,
)

# ========== SLIDE 2 STATUS ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_footer(slide, 2)
title_box(slide, "สถานะของโครงงาน")
tf = body_box(slide, 0.6, 1.2, 12, 5.4)
add_para(tf, "ชี้แจงก่อนนำเสนอ", 22, True, ACCENT, space_before=0)
for t in [
    "•  Digital Gift เป็นต้นแบบ (Prototype / แบบจำลองระบบ)",
    "•  ใช้สำหรับเรียนรู้ สาธิต และนำเสนอแนวคิดธุรกิจดิจิทัล",
    "•  มีหน้าเว็บตัวอย่างและระบบแอดมินให้ทดลองใช้งานได้",
    "•  สามารถสาธิตกระบวนการทำงานได้ครบวงจร",
    "•  ยังไม่ได้เปิดเป็นเว็บบริการเชิงพาณิชย์จริงในปัจจุบัน",
]:
    add_para(tf, t, 22, False, DARK, space_before=14)

# ========== SLIDE 3 BACKGROUND ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_footer(slide, 3)
title_box(slide, "ความเป็นมา")
items = [
    ("พฤติกรรมผู้บริโภค", "คนรุ่นใหม่ใช้มือถือและลิงก์ในการสื่อสารและมอบความรู้สึกมากขึ้น"),
    ("ช่องว่างของตลาด", "ของขวัญทั่วไปอาจไม่เฉพาะบุคคล และการส่งข้อความในแชทจางหายเร็ว"),
    ("แนวคิดต้นแบบ", "ออกแบบของขวัญดิจิทัลที่เล่นได้ ส่งด้วยลิงก์เดียว และปรับแต่งเฉพาะบุคคล"),
]
for i, (h, b) in enumerate(items):
    y = 1.3 + i * 1.7
    card(slide, 0.6, y, 12.1, 1.5)
    tf = body_box(slide, 0.9, y + 0.25, 11.5, 1.1)
    add_para(tf, f"{i+1}.  {h}", 22, True, NAVY, space_before=0)
    add_para(tf, b, 18, False, DARK, space_before=6)

# ========== SLIDE 4 PROBLEM OBJ ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_footer(slide, 4)
title_box(slide, "ปัญหาและวัตถุประสงค์")
card(slide, 0.6, 1.2, 5.9, 5.2)
card(slide, 6.8, 1.2, 5.9, 5.2)
tf = body_box(slide, 0.9, 1.4, 5.4, 4.8)
add_para(tf, "ปัญหาที่พบ", 24, True, ACCENT, space_before=0)
for t in [
    "1. ของขวัญทั่วไปขาดความเป็นส่วนตัว",
    "2. การส่งความรู้สึกผ่านแชทจางหายเร็ว",
    "3. ลูกค้าต้องการทดลองตัวอย่างก่อนตัดสินใจ",
    "4. ผู้ประกอบการรายย่อยทำเว็บเองได้ยาก",
]:
    add_para(tf, t, 18, False, DARK, space_before=16)
tf = body_box(slide, 7.1, 1.4, 5.4, 4.8)
add_para(tf, "วัตถุประสงค์", 24, True, ACCENT, space_before=0)
for t in [
    "1. พัฒนาต้นแบบเว็บของขวัญดิจิทัล",
    "2. ให้ทดลองเล่น Demo ก่อนสั่งทำ",
    "3. จำลองการรับออเดอร์ผ่าน LINE และการส่งมอบด้วยลิงก์",
    "4. ประยุกต์แนวคิดอีคอมเมิร์ซกับธุรกิจดิจิทัล",
]:
    add_para(tf, t, 18, False, DARK, space_before=16)

# ========== SLIDE 5 SCOPE ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_footer(slide, 5)
title_box(slide, "ขอบเขตของต้นแบบ")
scopes = [
    ("รูปแบบธุรกิจ", "B2C (ผู้บริโภครายบุคคล)"),
    ("ช่องทางสั่งซื้อ", "LINE (จำลองกระบวนการ)"),
    ("การส่งมอบ", "ลิงก์ของขวัญ + QR"),
    ("รูปแบบสินค้า", "งานสั่งทำตามเทมเพลต (Custom Order)"),
    ("สิ่งที่ไม่มี", "ไม่มีตะกร้าสินค้าแบบร้านค้าทั่วไป"),
    ("ราคา/ชำระเงิน", "เป็นแนวทางจำลองธุรกิจ ยังไม่เปิดขายจริง"),
]
for i, (h, b) in enumerate(scopes):
    col = i % 3
    row = i // 3
    x = 0.6 + col * 4.15
    y = 1.3 + row * 2.5
    card(slide, x, y, 3.95, 2.2)
    tf = body_box(slide, x + 0.25, y + 0.4, 3.45, 1.5)
    add_para(tf, h, 18, True, ACCENT, space_before=0)
    add_para(tf, b, 18, False, DARK, space_before=10)

# ========== SLIDE 6 BUSINESS ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_footer(slide, 6)
title_box(slide, "แนวคิดธุรกิจ (Business Model)")
rows = [
    ("คุณค่าที่มอบ", "ของขวัญดิจิทัลเฉพาะบุคคล ส่งผ่านลิงก์เดียว"),
    ("ลูกค้าเป้าหมาย", "วัยรุ่น–วัยทำงาน ที่ต้องการเซอร์ไพรส์คนสำคัญ"),
    ("ช่องทาง", "เว็บตัวอย่าง (Demo) + LINE"),
    ("รายได้ (แนวทาง)", "ค่าบริการจัดทำตามเทมเพลต"),
    ("ต้นทุนหลัก", "เวลาผลิต / โฮสติ้ง / การสื่อสารลูกค้า"),
]
for i, (h, b) in enumerate(rows):
    y = 1.25 + i * 0.95
    left = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(y), Inches(3.3), Inches(0.8)
    )
    left.fill.solid()
    left.fill.fore_color.rgb = NAVY
    left.line.fill.background()
    left.adjustments[0] = 0.15
    tf = body_box(slide, 0.75, y + 0.2, 3.0, 0.5)
    add_para(tf, h, 16, True, WHITE, space_before=0)
    right = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.1), Inches(y), Inches(8.6), Inches(0.8)
    )
    right.fill.solid()
    right.fill.fore_color.rgb = WHITE
    right.line.color.rgb = LIGHT
    right.adjustments[0] = 0.1
    tf = body_box(slide, 4.35, y + 0.2, 8.2, 0.5)
    add_para(tf, b, 18, False, DARK, space_before=0)

# ========== SLIDE 7 PRODUCTS ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_footer(slide, 7)
title_box(slide, "เทมเพลตของขวัญในต้นแบบ")
products = [
    ("กราดพุงเสืออวยพร", "59 บาท"),
    ("หน้ารำลึกความทรงจำ", "79 บาท"),
    ("เรื่องราวความรัก", "89 บาท"),
    ("ควิซความรักครบชุด", "99 บาท"),
    ("ของขวัญวันเกิด", "99 บาท"),
]
for i, (name, price) in enumerate(products):
    y = 1.2 + i * 0.85
    card(slide, 0.6, y, 12.1, 0.75)
    num = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(0.85), Inches(y + 0.15), Inches(0.45), Inches(0.45)
    )
    num.fill.solid()
    num.fill.fore_color.rgb = ACCENT
    num.line.fill.background()
    tf = body_box(slide, 0.85, y + 0.22, 0.45, 0.4)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(i + 1)
    set_run(run, 14, True, WHITE)
    tf = body_box(slide, 1.55, y + 0.18, 8, 0.45)
    add_para(tf, name, 20, True, NAVY, space_before=0)
    tf = body_box(slide, 10.2, y + 0.18, 2.2, 0.45)
    add_para(tf, price, 20, True, ACCENT, space_before=0)
tf = body_box(slide, 0.6, 5.55, 12, 0.8)
add_para(
    tf,
    "หมายเหตุ: ราคาเป็นราคาจำลองสำหรับนำเสนอแนวคิดธุรกิจ  •  ลูกค้าสามารถเล่น Demo ก่อนตัดสินใจได้",
    16,
    False,
    GRAY,
    space_before=0,
)

# ========== SLIDE 8 JOURNEY ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_footer(slide, 8)
title_box(slide, "Customer Journey ในแบบจำลอง")
steps = [
    ("1", "ทดลอง Demo", "เข้าเว็บเล่นตัวอย่าง"),
    ("2", "เลือกแบบ", "เลือกเทมเพลตที่ชอบ"),
    ("3", "ส่งข้อมูล", "ทัก LINE ส่งชื่อ ข้อความ รูป"),
    ("4", "จัดทำ", "สร้างของขวัญในแอดมิน"),
    ("5", "ส่งมอบ", "ชำระเงิน → ได้ลิงก์"),
]
for i, (n, h, b) in enumerate(steps):
    x = 0.45 + i * 2.55
    card(slide, x, 2.0, 2.4, 3.6)
    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x + 0.75), Inches(2.35), Inches(0.9), Inches(0.9)
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = NAVY
    circle.line.fill.background()
    tf = body_box(slide, x + 0.75, 2.5, 0.9, 0.6)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = n
    set_run(run, 24, True, WHITE)
    tf = body_box(slide, x + 0.15, 3.5, 2.1, 1.8)
    add_para(tf, h, 18, True, ACCENT, PP_ALIGN.CENTER, 0)
    add_para(tf, b, 15, False, DARK, PP_ALIGN.CENTER, 10)
    if i < 4:
        arrow = slide.shapes.add_shape(
            MSO_SHAPE.RIGHT_ARROW, Inches(x + 2.35), Inches(3.5), Inches(0.28), Inches(0.28)
        )
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = ACCENT
        arrow.line.fill.background()

# ========== SLIDE 9 SYSTEM ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_footer(slide, 9)
title_box(slide, "ระบบและการทำงานของต้นแบบ")
card(slide, 0.6, 1.3, 6.0, 5.1)
card(slide, 6.9, 1.3, 5.8, 5.1)
tf = body_box(slide, 0.9, 1.55, 5.5, 4.6)
add_para(tf, "ฝั่งลูกค้า", 24, True, ACCENT, space_before=0)
for t in [
    "•  หน้าแรกแนะนำบริการและตัวอย่าง",
    "•  หน้า Demo เล่นเทมเพลตได้ทันที",
    "•  หน้าของขวัญจริงจากลิงก์เฉพาะ",
    "•  รองรับการใช้งานบนมือถือ",
]:
    add_para(tf, t, 18, False, DARK, space_before=18)
tf = body_box(slide, 7.2, 1.55, 5.3, 4.6)
add_para(tf, "ฝั่งแอดมิน", 24, True, ACCENT, space_before=0)
for t in [
    "•  เข้าสู่ระบบด้วยบัญชีผู้ดูแล",
    "•  สร้าง / แก้ไข / จัดการของขวัญ",
    "•  ใส่ชื่อ ข้อความ และรูปตามเทมเพลต",
    "•  เผยแพร่และส่งมอบด้วยลิงก์",
]:
    add_para(tf, t, 18, False, DARK, space_before=18)

# ========== SLIDE 10 TECH ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_footer(slide, 10)
title_box(slide, "เทคโนโลยีที่ใช้ (Tech Stack)")
techs = [
    ("Frontend", "React + TypeScript + Vite"),
    ("Backend", "Go (Chi) + JWT"),
    ("Database", "PostgreSQL"),
    ("ช่องทางสื่อสาร", "LINE"),
    ("ตัวอย่างสาธารณะ", "GitHub Pages (สำหรับสาธิต)"),
]
for i, (h, b) in enumerate(techs):
    y = 1.25 + i * 0.95
    card(slide, 0.6, y, 12.1, 0.85)
    tf = body_box(slide, 0.95, y + 0.22, 3.5, 0.5)
    add_para(tf, h, 20, True, ACCENT, space_before=0)
    tf = body_box(slide, 4.5, y + 0.22, 7.8, 0.5)
    add_para(tf, b, 20, False, DARK, space_before=0)

# ========== SLIDE 11 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_footer(slide, 11)
title_box(slide, "จุดเด่น ข้อจำกัด และบทเรียน")
cols = [
    (
        "จุดเด่น",
        [
            "เล่น Demo ได้ก่อนตัดสินใจ",
            "ของขวัญเฉพาะบุคคล ส่งด้วยลิงก์",
            "มีหลายเทมเพลต",
            "สาธิตกระบวนการธุรกิจได้ครบ",
        ],
        NAVY,
    ),
    (
        "ข้อจำกัด",
        [
            "ยังไม่เปิดขายจริง",
            "ชำระเงินยังอยู่นอกระบบ",
            "บางขั้นอาศัยสื่อสารผ่าน LINE",
            "ยังเป็นต้นแบบเพื่อการเรียนรู้",
        ],
        ACCENT,
    ),
    (
        "บทเรียน",
        [
            "การสื่อสารลูกค้าต้องชัดเป็นขั้น",
            "Demo ช่วยให้เข้าใจสินค้าเร็วขึ้น",
            "โครงสร้างระบบควรแยกฝั่งชัดเจน",
            "ธุรกิจดิจิทัลต้องครบทั้งประสบการณ์และกระบวนการ",
        ],
        NAVY,
    ),
]
for i, (h, items, color) in enumerate(cols):
    x = 0.5 + i * 4.2
    card(slide, x, 1.25, 4.0, 5.2)
    head = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.25), Inches(4.0), Inches(0.7)
    )
    head.fill.solid()
    head.fill.fore_color.rgb = color
    head.line.fill.background()
    tf = body_box(slide, x + 0.2, 1.38, 3.6, 0.5)
    add_para(tf, h, 20, True, WHITE, PP_ALIGN.CENTER, 0)
    tf = body_box(slide, x + 0.25, 2.2, 3.5, 4.0)
    for t in items:
        add_para(tf, "•  " + t, 16, False, DARK, space_before=14)

# ========== SLIDE 12 SUMMARY ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_footer(slide, 12)
title_box(slide, "สรุปและแนวทางต่อยอด")
card(slide, 0.6, 1.2, 6.0, 4.6)
card(slide, 6.9, 1.2, 5.8, 4.6)
tf = body_box(slide, 0.9, 1.45, 5.5, 4.2)
add_para(tf, "สรุป", 24, True, ACCENT, space_before=0)
for t in [
    "•  พัฒนาต้นแบบ Digital Gift ได้สำเร็จในระดับสาธิต",
    "•  แสดงแนวคิดอีคอมเมิร์ซแบบ B2C + Custom Fulfillment",
    "•  ได้เรียนรู้ทั้งธุรกิจ การสื่อสารลูกค้า และการพัฒนาเว็บ",
    "•  ยังไม่ใช่เว็บบริการจริง แต่พร้อมเป็นฐานต่อยอด",
]:
    add_para(tf, t, 17, False, DARK, space_before=16)
tf = body_box(slide, 7.2, 1.45, 5.3, 4.2)
add_para(tf, "แนวทางต่อยอด", 24, True, ACCENT, space_before=0)
for t in [
    "•  พัฒนาสู่เว็บบริการจริง",
    "•  เพิ่มระบบชำระเงินในระบบ",
    "•  ขยายเทมเพลตและช่องทางการตลาด",
    "•  ปรับปรุงประสบการณ์มือถือให้สมบูรณ์ขึ้น",
]:
    add_para(tf, t, 17, False, DARK, space_before=16)

thanks = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(6.0), Inches(12.1), Inches(0.7)
)
thanks.fill.solid()
thanks.fill.fore_color.rgb = NAVY
thanks.line.fill.background()
thanks.adjustments[0] = 0.2
tf = body_box(slide, 0.6, 6.15, 12.1, 0.5)
add_para(tf, "ขอบคุณครับ/ค่ะ ที่รับฟังการนำเสนอ", 20, True, WHITE, PP_ALIGN.CENTER, 0)

out = r"c:\Users\User\OneDrive\Desktop\Digital Gift\docs\Digital_Gift_Presentation.pptx"
os.makedirs(os.path.dirname(out), exist_ok=True)
prs.save(out)
print("SAVED", out)
